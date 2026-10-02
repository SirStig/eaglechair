"""
Catalog Builder, product exports and product register (admin)

Covers the saved-project CRUD, page previews and PDF export in the catalog
design, the Excel / product index downloads and the register's bulk edits.
"""

import base64
import io

import fitz
import pytest
from openpyxl import load_workbook
from PIL import Image

from backend.models.chair import Chair, ProductVariation
from backend.models.company import AdminRole, AdminUser
from tests.factories import create_category, create_chair, create_product_family, create_product_variation

UA = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) test", "X-Session-Token": "s", "X-Admin-Token": "a"}
BASE = "/api/v1/admin"


@pytest.fixture
def as_role():
    from backend.api.dependencies import get_current_admin
    from tests.conftest import get_app

    app = get_app()

    def _set(role=AdminRole.ADMIN):
        admin = AdminUser(id=525252, username=f"catalog-{role.value}", email="catalog@example.com",
                          role=role, is_active=True)
        app.dependency_overrides[get_current_admin] = lambda: admin

    yield _set
    app.dependency_overrides.pop(get_current_admin, None)


@pytest.fixture
def uploads(tmp_path, monkeypatch):
    """A transparent product cut-out under a temporary uploads/images folder."""
    from backend.api.v1.routes.admin import catalog_builder

    folder = tmp_path / "images" / "products"
    folder.mkdir(parents=True)
    img = Image.new("RGBA", (400, 600), (0, 0, 0, 0))
    img.paste((200, 60, 30, 255), (100, 100, 300, 560))
    img.save(folder / "chair.png")
    monkeypatch.setattr(catalog_builder, "get_upload_base_dir", lambda: tmp_path)
    return "/uploads/images/products/chair.png"


@pytest.fixture
async def lobo(db_session, uploads):
    category = await create_category(db_session, spec_profile="chair")
    family = await create_product_family(db_session, category_id=category.id, name="Lobo", overview_text="Modern design.")
    side = await create_chair(
        db_session, category_id=category.id, family_id=family.id, model_number="3506", model_suffix="CR",
        name="Lobo Side Chair", height=32, width=23.5, seat_height=18, images=[{"url": uploads}],
        frame_material="Steel frame",
    )
    stool = await create_chair(
        db_session, category_id=category.id, family_id=family.id, model_number="4506", name="Lobo Barstool",
        is_active=False,
    )
    variations = [
        await create_product_variation(
            db_session, side.id, sku=f"3506.{code}", name=f"{code} frame", images=[uploads], display_order=order
        )
        for order, code in enumerate(("Bl", "Cr", "Gd"))
    ]
    return {"family": family, "side": side, "stool": stool, "variations": variations}


def _document(lobo):
    return {
        "settings": {"title": "Test Catalog"},
        "pages": [
            {"type": "cover", "title": "Indoor", "items": [{"product_id": lobo["side"].id}]},
            {"type": "toc"},
            {"type": "product", "title": "Lobo", "items": [{"product_id": lobo["side"].id}, {"product_id": lobo["stool"].id}],
             "features": "Modern.", "emblem": "flag"},
            {"type": "gallery", "title": "Lobo", "subtitle": "variations",
             "items": [{"variation_id": v.id} for v in lobo["variations"]]},
            {"type": "photo", "image_url": "/uploads/images/products/chair.png", "include_in_toc": False},
        ],
    }


@pytest.mark.integration
@pytest.mark.admin
class TestCatalogBuilder:
    async def test_project_crud(self, async_client, as_role):
        as_role(AdminRole.EDITOR)
        created = await async_client.post(f"{BASE}/catalog-builder/projects", json={"name": "Spring 2026"}, headers=UA)
        assert created.status_code == 201, created.text
        project = created.json()
        assert project["document"]["pages"] == []

        doc = {"pages": [{"type": "toc", "title": "Contents"}]}
        saved = await async_client.put(f"{BASE}/catalog-builder/projects/{project['id']}", json={"document": doc}, headers=UA)
        assert saved.status_code == 200
        assert saved.json()["page_count"] == 1
        assert saved.json()["document"]["pages"][0]["id"]  # ids are assigned

        listed = await async_client.get(f"{BASE}/catalog-builder/projects", headers=UA)
        assert [p["name"] for p in listed.json()] == ["Spring 2026"]

        copy = await async_client.post(f"{BASE}/catalog-builder/projects/{project['id']}/duplicate", headers=UA)
        assert copy.json()["name"] == "Spring 2026 (copy)"

        # Editors cannot delete; admins can
        denied = await async_client.delete(f"{BASE}/catalog-builder/projects/{project['id']}", headers=UA)
        assert denied.status_code == 403
        as_role(AdminRole.ADMIN)
        deleted = await async_client.delete(f"{BASE}/catalog-builder/projects/{project['id']}", headers=UA)
        assert deleted.status_code == 200
        missing = await async_client.get(f"{BASE}/catalog-builder/projects/{project['id']}", headers=UA)
        assert missing.status_code == 404

    async def test_viewer_cannot_create_projects(self, async_client, as_role):
        as_role(AdminRole.VIEWER)
        response = await async_client.post(f"{BASE}/catalog-builder/projects", json={"name": "x"}, headers=UA)
        assert response.status_code == 403

    async def test_preview_returns_png_and_photo_slots(self, async_client, as_role, lobo):
        as_role()
        response = await async_client.post(
            f"{BASE}/catalog-builder/preview", json={"document": _document(lobo), "page_index": 2}, headers=UA
        )
        assert response.status_code == 200, response.text
        body = response.json()
        png = base64.b64decode(body["image"].split(",", 1)[1])
        assert png.startswith(b"\x89PNG")
        assert body["page_number"] == 3
        assert body["total_pages"] == 5
        # Side chair has a photo, the barstool does not
        assert [s["missing"] for s in body["slots"]] == [False, True]

    async def test_preview_rejects_bad_page_index(self, async_client, as_role, lobo):
        as_role()
        response = await async_client.post(
            f"{BASE}/catalog-builder/preview", json={"document": _document(lobo), "page_index": 9}, headers=UA
        )
        assert response.status_code == 400

    async def test_export_pdf(self, async_client, as_role, lobo):
        as_role()
        response = await async_client.post(
            f"{BASE}/catalog-builder/export", json={"document": _document(lobo), "filename": "Spring/../2026"}, headers=UA
        )
        assert response.status_code == 200, response.text
        assert response.headers["content-type"] == "application/pdf"
        assert 'filename="Spring..2026.pdf"' in response.headers["content-disposition"]
        pdf = fitz.open("pdf", response.content)
        assert len(pdf) == 5
        assert "Lobo" in pdf[1].get_text()  # contents
        product_text = pdf[2].get_text()
        assert "3506 CR" in product_text and "32”" in product_text and "Features" in product_text
        assert "3506.Bl" in pdf[3].get_text()  # gallery captions

    async def test_suggest_pages_for_family(self, async_client, as_role, lobo):
        as_role()
        response = await async_client.post(
            f"{BASE}/catalog-builder/suggest-pages", json={"family_ids": [lobo["family"].id]}, headers=UA
        )
        assert response.status_code == 200
        pages = response.json()["pages"]
        assert [p["type"] for p in pages] == ["product", "gallery"]
        assert pages[0]["title"] == "Lobo"
        assert pages[0]["features"] == "Modern design."  # family overview
        assert pages[0]["materials"] == "Steel frame"
        assert [i["product_id"] for i in pages[0]["items"]] == [lobo["side"].id]  # inactive barstool left out
        assert len(pages[1]["items"]) == 3

    async def test_picker_products_search(self, async_client, as_role, lobo):
        as_role()
        response = await async_client.get(f"{BASE}/catalog-builder/products", params={"search": "3506"}, headers=UA)
        products = response.json()
        assert [p["model_number"] for p in products] == ["3506"]
        assert [v["sku"] for v in products[0]["variations"]] == ["3506.Bl", "3506.Cr", "3506.Gd"]


@pytest.mark.integration
@pytest.mark.admin
class TestProductExports:
    async def test_excel_export(self, async_client, as_role, lobo):
        as_role(AdminRole.VIEWER)
        response = await async_client.get(f"{BASE}/exports/products.xlsx", headers=UA)
        assert response.status_code == 200, response.text
        wb = load_workbook(io.BytesIO(response.content))
        products = list(wb["Products"].iter_rows(values_only=True))
        assert products[0][:4] == ("Model Number", "Suffix", "Full Model", "Name")
        rows = {r[2]: r for r in products[1:]}
        assert rows["3506 CR"][3] == "Lobo Side Chair"
        assert rows["3506 CR"][10] == "3506.Bl, 3506.Cr, 3506.Gd"
        assert rows["4506"][7] == "Inactive"
        variations = list(wb["Variations"].iter_rows(values_only=True))
        assert [r[3] for r in variations[1:]] == ["3506.Bl", "3506.Cr", "3506.Gd"]

    async def test_excel_export_active_only(self, async_client, as_role, lobo):
        as_role()
        response = await async_client.get(f"{BASE}/exports/products.xlsx", params={"include_inactive": False}, headers=UA)
        wb = load_workbook(io.BytesIO(response.content))
        assert [r[2] for r in wb["Products"].iter_rows(min_row=2, values_only=True)] == ["3506 CR"]

    async def test_product_index_pdf(self, async_client, as_role, lobo):
        as_role()
        response = await async_client.get(f"{BASE}/exports/product-index.pdf", params={"include_inactive": True}, headers=UA)
        assert response.status_code == 200
        text = fitz.open("pdf", response.content)[0].get_text()
        assert "Product Index" in text
        assert "Lobo Side Chair" in text and "3506.Bl" in text
        assert "(inactive)" in text


@pytest.mark.integration
@pytest.mark.admin
class TestProductRegister:
    async def test_register_lists_everything_with_issues(self, async_client, as_role, lobo):
        as_role(AdminRole.VIEWER)
        response = await async_client.get(f"{BASE}/register", headers=UA)
        assert response.status_code == 200
        rows = {r["model_number"]: r for r in response.json()["products"]}
        assert set(rows) == {"3506", "4506"}
        assert "no_photo" not in rows["3506"]["issues"]
        assert "no_photo" in rows["4506"]["issues"]
        assert rows["4506"]["is_active"] is False
        assert len(rows["3506"]["variations"]) == 3

    async def test_bulk_update(self, async_client, as_role, lobo, db_session):
        as_role()
        ids = [lobo["side"].id, lobo["stool"].id]
        response = await async_client.post(f"{BASE}/register/bulk", json={"product_ids": ids, "is_active": True}, headers=UA)
        assert response.status_code == 200, response.text
        assert response.json() == {"updated": 2}
        for product_id in ids:
            product = await db_session.get(Chair, product_id)
            await db_session.refresh(product)
            assert product.is_active is True

    async def test_bulk_update_requires_admin(self, async_client, as_role, lobo):
        as_role(AdminRole.EDITOR)
        response = await async_client.post(
            f"{BASE}/register/bulk", json={"product_ids": [lobo["side"].id], "is_active": False}, headers=UA
        )
        assert response.status_code == 403

    async def test_variation_patch_and_sku_clash(self, async_client, as_role, lobo, db_session):
        as_role()
        first, second = lobo["variations"][:2]
        clash = await async_client.patch(f"{BASE}/register/variations/{first.id}", json={"sku": second.sku}, headers=UA)
        assert clash.status_code == 409
        ok = await async_client.patch(
            f"{BASE}/register/variations/{first.id}", json={"sku": "3506.BL2", "is_available": False}, headers=UA
        )
        assert ok.status_code == 200, ok.text
        variation = await db_session.get(ProductVariation, first.id)
        await db_session.refresh(variation)
        assert (variation.sku, variation.is_available) == ("3506.BL2", False)
